"""Schreibt aus den Markdown-Tokens eines Handouts eine Word-Datei zum Weiterbearbeiten.

Bewusst schlicht: Standardschrift (Calibri), grüne Überschriften, echte Nummerierung und
Aufzählungen, Auswahl und Hinweiskästen als farbige Tabellenzellen. Nur-Web-Blöcke
(::: weiterlesen) und Nur-Druck-Blöcke (::: nurdruck) gehören nicht in die Word-Datei.
"""

import re

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import _Cell

GRUEN = RGBColor(0x00, 0x66, 0x33)
ORANGE = "E3780C"
HELL = "EDEDDF"
ASCHE = "F7F7F1"
NICHT_IN_WORD = ("container_weiterlesen_open", "container_nurdruck_open")


def _fuellen(zelle, farbe):
    tc_pr = zelle._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), farbe)
    tc_pr.append(shd)


def _rahmen(tabelle, farbe, breite=18):
    """Äußerer Rahmen und Trennlinien in einer Farbe (Breite in 1/8 pt)."""
    tbl_pr = tabelle._tbl.tblPr
    rahmen = OxmlElement("w:tblBorders")
    for seite in ("top", "left", "bottom", "right", "insideH"):
        kante = OxmlElement(f"w:{seite}")
        kante.set(qn("w:val"), "single")
        kante.set(qn("w:sz"), str(breite))
        kante.set(qn("w:color"), farbe)
        rahmen.append(kante)
    tbl_pr.append(rahmen)


def _abstand_nach_tabelle(doc):
    """Word klebt zwei Tabellen sonst zusammen und lässt keinen Abstand zum Text danach."""
    absatz = doc.add_paragraph()
    absatz.paragraph_format.space_after = Pt(2)


class WordBauer:
    def __init__(self, daten):
        self.doc = Document()
        self.daten = daten
        self.ziel = self.doc
        self.listen = []  # Stapel aus "ol" / "ul"
        self.in_item_erster = []  # pro Listenebene: kommt noch der erste Absatz des Eintrags?
        self.wahl = None  # Tabelle der gerade offenen Auswahl
        self.in_variante = False
        self.wahl_intro = None  # Zelle mit der Erklärung unter dem Kopf der Auswahl
        self._stile()
        self.doc.core_properties.title = daten["titel"]
        self.doc.core_properties.author = "Julia Junge"

    # ---------- Stile ----------
    def _stile(self):
        sections = self.doc.sections
        for s in sections:
            s.left_margin = s.right_margin = Cm(2.2)
            s.top_margin = s.bottom_margin = Cm(2)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, groesse in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
            stil = self.doc.styles[name]
            stil.font.name = "Calibri"
            stil.font.size = Pt(groesse)
            stil.font.bold = True
            stil.font.color.rgb = GRUEN
            rfonts = stil.element.rPr.rFonts
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rfonts.get(qn(attr)) is not None:
                    del rfonts.attrib[qn(attr)]
            rfonts.set(qn("w:ascii"), "Calibri")
            rfonts.set(qn("w:hAnsi"), "Calibri")
        self.doc.styles["Heading 1"].paragraph_format.space_before = Pt(18)
        self.doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
        self.doc.styles["Heading 2"].paragraph_format.space_before = Pt(12)

    # ---------- Bausteine ----------
    def _absatz(self, stil=None):
        ziel = self.ziel
        if isinstance(ziel, _Cell):
            erster = ziel.paragraphs[0]
            if not erster.runs and not getattr(ziel, "_benutzt", False):
                ziel._benutzt = True
                if stil:
                    erster.style = self.doc.styles[stil]
                return erster
            return ziel.add_paragraph(style=stil)
        return ziel.add_paragraph(style=stil)

    def _inline(self, absatz, kinder, fett=False):
        fett_zaehler = 1 if fett else 0
        kursiv = 0
        for k in kinder:
            if k.type == "strong_open":
                fett_zaehler += 1
            elif k.type == "strong_close":
                fett_zaehler -= 1
            elif k.type == "em_open":
                kursiv += 1
            elif k.type == "em_close":
                kursiv -= 1
            elif k.type == "softbreak":
                absatz.add_run(" ")
            elif k.type == "hardbreak":
                absatz.add_run().add_break()
            elif k.type in ("text", "code_inline"):
                # ==Text== ist im Handout die gelbe Marker-Hervorhebung
                teile = re.split(r"==(.+?)==", k.content)
                for n, teil in enumerate(teile):
                    if not teil:
                        continue
                    lauf = absatz.add_run(teil)
                    lauf.bold = fett_zaehler > 0
                    lauf.italic = kursiv > 0
                    if n % 2 == 1:
                        lauf.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def _kasten(self, farbe, rahmen=None):
        """Einzellige Tabelle als Hinweiskasten; liefert die Zelle."""
        tabelle = self.doc.add_table(rows=1, cols=1)
        if rahmen:
            _rahmen(tabelle, rahmen)
        zelle = tabelle.rows[0].cells[0]
        _fuellen(zelle, farbe)
        return zelle

    def _wahl_zeile(self, farbe):
        zeile = self.wahl.add_row()
        _fuellen(zeile.cells[0], farbe)
        return zeile.cells[0]

    def _label(self, text):
        absatz = self._absatz()
        lauf = absatz.add_run(text.upper())
        lauf.bold = True
        lauf.font.size = Pt(9)
        return absatz

    # ---------- Durchlauf ----------
    def bauen(self, tokens):
        self.doc.add_heading(self.daten["titel"], level=0)
        i = 0
        while i < len(tokens):
            t = tokens[i]
            if t.type in NICHT_IN_WORD:
                # bis zum passenden Ende überspringen
                ende = t.type.replace("_open", "_close")
                while tokens[i].type != ende:
                    i += 1
            elif t.type == "heading_open":
                stufe = 1 if t.tag == "h2" else 2
                self.doc.add_heading(tokens[i + 1].content, level=stufe)
                i += 2
            elif t.type == "paragraph_open":
                self._paragraph(tokens[i + 1])
                i += 1
            elif t.type in ("ordered_list_open", "bullet_list_open"):
                self.listen.append("ol" if t.type.startswith("ordered") else "ul")
                self.in_item_erster.append(False)
            elif t.type in ("ordered_list_close", "bullet_list_close"):
                self.listen.pop()
                self.in_item_erster.pop()
            elif t.type == "list_item_open":
                self.in_item_erster[-1] = True
            elif t.type == "container_karte_open":
                label = t.info.strip()[len("karte"):].strip()
                self.ziel = self._kasten(HELL)
                if label:
                    self._label(label)
            elif t.type == "container_karte_close":
                self.ziel = self.doc
                _abstand_nach_tabelle(self.doc)
            elif t.type == "container_wahl_open":
                titel = t.info.strip()[len("wahl"):].strip()
                self.wahl_intro = None
                self.wahl = self.doc.add_table(rows=1, cols=1)
                _rahmen(self.wahl, ORANGE, 24)
                kopf = self.wahl.rows[0].cells[0]
                _fuellen(kopf, ORANGE)
                self.ziel = kopf
                self._label("Entscheide dich")
                lauf = self._absatz().add_run(titel)
                lauf.bold = True
            elif t.type == "container_wahl_close":
                self.wahl = None
                self.ziel = self.doc
                _abstand_nach_tabelle(self.doc)
            elif t.type == "container_variante_open":
                label = t.info.strip()[len("variante"):].strip()
                self.ziel = self._wahl_zeile(ASCHE)
                self.in_variante = True
                absatz = self._absatz()
                absatz.add_run("☐  " + label).bold = True
            elif t.type == "container_variante_close":
                self.in_variante = False
                self.ziel = self.doc
            i += 1
        self._schluss()

    def _paragraph(self, inline):
        if self.wahl is not None and not self.in_variante:
            # Erklärung direkt unter dem Orange-Kopf der Auswahl
            if self.wahl_intro is None:
                self.wahl_intro = self._wahl_zeile("FFFFFF")
            self.ziel = self.wahl_intro
        stil = None
        einzug = False
        if self.listen:
            erste = self.in_item_erster[-1]
            tiefe = len(self.listen)
            if erste:
                stil = "List Number" if self.listen[-1] == "ol" else ("List Bullet" if tiefe == 1 else "List Bullet 2")
                self.in_item_erster[-1] = False
            else:
                einzug = True
        absatz = self._absatz(stil)
        if einzug:
            absatz.paragraph_format.left_indent = Cm(0.63)
        self._inline(absatz, inline.children or [])

    def _schluss(self):
        absatz = self.doc.add_paragraph()
        absatz.paragraph_format.space_before = Pt(18)
        quelle = self.daten.get("quelle") or f"Autorin: Julia Junge · Stand: {self.daten['stand']}"
        lauf = absatz.add_run(quelle)
        lauf.font.size = Pt(9)
        lauf.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def schreibe_word(tokens, daten, ziel):
    bauer = WordBauer(daten)
    bauer.bauen(tokens)
    bauer.doc.save(ziel)
